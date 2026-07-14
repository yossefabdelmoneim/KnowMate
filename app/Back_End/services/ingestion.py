import os
import uuid
from langchain_community.document_loaders import PyPDFLoader, TextLoader, Docx2txtLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.Back_End.db.vector_store import get_vector_store


def load_excel(path: str) -> list[Document]:
    if path.lower().endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(path)
        docs = []
        for sheet in wb.sheets():
            rows = []
            for row_idx in range(sheet.nrows):
                cells = [str(c) for c in sheet.row_values(row_idx) if c]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                text = "\n".join(rows)
                docs.append(Document(page_content=text, metadata={"source": os.path.basename(path)}))
        return docs
    from openpyxl import load_workbook
    wb = load_workbook(path)
    docs = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            text = "\n".join(rows)
            docs.append(Document(page_content=text, metadata={"source": os.path.basename(path)}))
    wb.close()
    return docs


def load_csv(path: str) -> list[Document]:
    import csv
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = []
        for row in reader:
            if row:
                rows.append(" | ".join(row))
    if not rows:
        return []
    text = "\n".join(rows)
    return [Document(page_content=text, metadata={"source": os.path.basename(path)})]


def load_document(path: str):
    #Detects file type and then load it
    suffix = path.lower()

    if suffix.endswith(".pdf"):
        return PyPDFLoader(path).load()
    elif suffix.endswith(".docx"):
        return Docx2txtLoader(path).load()
    elif suffix.endswith(".xlsx") or suffix.endswith(".xls"):
        return load_excel(path)
    elif suffix.endswith(".csv"):
        return load_csv(path)
    elif suffix.endswith(".txt"):
        return TextLoader(path).load()

    raise ValueError(f"Unsupported file type: {path}")


def split_docs(docs):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=100
    )
    return splitter.split_documents(docs)


def add_metadata(chunks, company_id, user_id=None, doc_id=None, source=None):
    doc_id = doc_id or str(uuid.uuid4()) # Generate a unique document ID

    for i, c in enumerate(chunks):
        meta = {
            "doc_id": doc_id,
            "chunk_id": i,
            "company_id": company_id,
            "source": source or "",
        }
        if user_id is not None:
            meta["user_id"] = user_id
        c.metadata.update(meta)

    return chunks


def build_chunks(path: str, company_id: str, user_id: str | None = None, doc_id: str | None = None, source: str | None = None):
    docs = load_document(path)
    chunks = split_docs(docs)

    if not chunks:
        raise ValueError("Document did not contain any readable text.")

    return add_metadata(chunks, company_id, user_id, doc_id, source)


def add_chunks(chunks):
    db = get_vector_store()
    db.add_documents(chunks)

    return {
        "doc_id": chunks[0].metadata["doc_id"],
        "chunks": len(chunks)
    }


def ingest(path: str, company_id: str, user_id: str | None = None, doc_id: str | None = None, source: str | None = None):
    chunks = build_chunks(path, company_id, user_id, doc_id, source)
    return add_chunks(chunks)


def get_document_chunk_ids(doc_id: str, company_id: str | None = None):
    db = get_vector_store()
    where = {"doc_id": doc_id}
    result = db.get(where=where, include=["metadatas"])

    ids = result.get("ids", [])
    metadatas = result.get("metadatas", [])

    if company_id is not None:
        ids = [
            chunk_id
            for chunk_id, metadata in zip(ids, metadatas)
            if metadata.get("company_id") == company_id
        ]

    return ids


def delete_chunk_ids(ids):
    if not ids:
        return 0

    db = get_vector_store()
    db.delete(ids=ids)
    return len(ids)


def delete_document(doc_id: str, company_id: str | None = None):
    ids = get_document_chunk_ids(doc_id, company_id)
    return delete_chunk_ids(ids)


def replace_document(path: str, company_id: str, doc_id: str, source: str | None = None):
    old_ids = get_document_chunk_ids(doc_id, company_id)

    if not old_ids:
        return None

    chunks = build_chunks(path, company_id, doc_id, source)
    result = add_chunks(chunks)
    result["replaced_chunks"] = delete_chunk_ids(old_ids)

    return result
