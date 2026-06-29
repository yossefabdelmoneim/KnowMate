from openpyxl import load_workbook


def load_excel(file_path: str) -> str:
    """
    Read an Excel file and return its content as text.
    """

    workbook = load_workbook(file_path)

    text = ""

    for sheet in workbook.worksheets:

        text += f"\n===== Sheet: {sheet.title} =====\n"

        for row in sheet.iter_rows(values_only=True):

            row_text = []

            for cell in row:

                if cell is not None:
                    row_text.append(str(cell))

            if row_text:
                text += " | ".join(row_text) + "\n"

    return text